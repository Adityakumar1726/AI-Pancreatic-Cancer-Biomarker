from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

root = Path(".")
results_dir = root / "results"

csv_path = results_dir / "INDEPENDENT_BIOMARKER_VALIDATION_CORRECTED.csv"
plot_path = results_dir / "INDEPENDENT_VALIDATION_MROH6_OSCAR.png"
report_path = results_dir / "INDEPENDENT_VALIDATION_REPORT.docx"

df = pd.read_csv(csv_path)

doc = Document()
doc.add_heading("Independent Validation of Pancreatic Cancer Biomarker Candidates", 0)
doc.add_paragraph(
    "Project: AI-Based Pancreatic Cancer Biomarker Identification "
    "Using Relational Graph Convolutional Networks and Deep Q-Networks"
)

doc.add_heading("1. Objective", level=1)
doc.add_paragraph(
    "To examine whether selected computationally prioritized biomarker candidates "
    "show expression differences between pancreatic tumor and adjacent non-tumor "
    "samples in an independent public dataset."
)

doc.add_heading("2. Dataset", level=1)
doc.add_paragraph(
    "Dataset: GSE183795. The comparison included 139 tumor samples and "
    "102 adjacent non-tumor samples. Three normal-pancreas donor samples were "
    "excluded from this two-group comparison."
)

doc.add_heading("3. Methodology", level=1)
doc.add_paragraph(
    "Normalized microarray expression values were compared between tumor and "
    "adjacent non-tumor groups. Welch's two-sample t-test was used for each "
    "candidate, followed by Benjamini-Hochberg false discovery rate (FDR) "
    "correction across the successfully tested candidates. An FDR below 0.05 "
    "was considered statistically significant for this exploratory analysis."
)

doc.add_heading("4. Results", level=1)
table = doc.add_table(rows=1, cols=6)
table.style = "Light Shading Accent 1"
headers = [
    "Gene", "Tumor mean", "Adjacent mean",
    "Mean difference", "P-value", "FDR"
]
for cell, value in zip(table.rows[0].cells, headers):
    cell.text = value

for _, row in df.iterrows():
    values = [
        str(row["gene_symbol"]),
        f'{row["tumor_mean"]:.4f}' if pd.notna(row["tumor_mean"]) else "NA",
        f'{row["adjacent_mean"]:.4f}' if pd.notna(row["adjacent_mean"]) else "NA",
        f'{row["mean_difference"]:.4f}' if pd.notna(row["mean_difference"]) else "NA",
        f'{row["p_value"]:.3g}' if pd.notna(row["p_value"]) else "NA",
        f'{row["FDR"]:.3g}' if pd.notna(row.get("FDR")) else "NA"
    ]
    cells = table.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value

doc.add_heading("5. Expression Plot", level=1)
if plot_path.exists():
    doc.add_picture(str(plot_path), width=Inches(6.2))
    doc.add_paragraph(
        "Figure 1. Normalized expression distributions for MROH6 and OSCAR "
        "in tumor and adjacent non-tumor samples."
    )
else:
    doc.add_paragraph("The expression plot was not found.")

doc.add_heading("6. Interpretation", level=1)
doc.add_paragraph(
    "MROH6 showed higher average expression in tumor samples, with an FDR of "
    "approximately 4.38 × 10⁻¹². OSCAR showed lower average expression in tumor "
    "samples, with an FDR of approximately 1.55 × 10⁻⁴. Both met the exploratory "
    "FDR threshold of 0.05. The other successfully tested candidates did not meet "
    "that threshold in this comparison."
)

doc.add_heading("7. Limitations", level=1)
doc.add_paragraph(
    "This analysis evaluates expression differences, not diagnostic performance "
    "or clinical utility. The microarray may not represent every transcript "
    "or candidate. FCGR1BP was not matched by its official symbol in the matrix. "
    "The statistical analysis does not account for all possible confounders, "
    "such as patient pairing, batch effects, or clinical characteristics. "
    "Independent replication and further biological and clinical assessment "
    "are needed before any candidate can be considered clinically validated."
)

doc.add_heading("8. Conclusion", level=1)
doc.add_paragraph(
    "In this exploratory independent dataset analysis, MROH6 and OSCAR showed "
    "statistically significant tumor-versus-adjacent-tissue expression "
    "differences after FDR correction. These results support further study "
    "but do not establish that either gene is a clinically validated "
    "pancreatic cancer biomarker."
)

doc.save(report_path)
print("Report created:", report_path)
print("Candidates included:", len(df))
